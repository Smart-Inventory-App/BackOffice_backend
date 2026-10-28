"""
`app/api/me.py`: perfil, filtros, estadísticas, listado/exportación de
productos, edición en lote y carga masiva (CSV / Excel).

Lo específico del código de barras ya está cubierto en `test_barcode.py`; acá
se cubre el resto de las ramas de `_validate_row`, `_coerce_*`, `bulk_parse` y
`bulk_commit`, más los endpoints de perfil/estadísticas/listado que todavía no
tenían test propio.
"""
import io
from datetime import date, datetime, timedelta

from app.core.argentina import ar_today

import pytest
from openpyxl import Workbook
from sqlalchemy.exc import SQLAlchemyError

from app.api import me as me_api
from app.models.app_models import ConsumptionEvent, Product, UserProfile
from app.models.audit import AdminActivityLog
from tests.conftest import create_product, create_stock_alert

CSV_HEADER = "nombre,categoria,stock,codigo_barras,vencimiento,marca,ubicacion,descripcion,imagen_url,stock_minimo,dias_alerta_vencimiento\n"


def _csv_file(contenido: str):
    return ("productos.csv", contenido.encode("utf-8"), "text/csv")


def _parse(client, contenido: str, mode: str = "entrada"):
    resp = client.post(
        "/me/products/bulk/parse", params={"mode": mode}, files={"file": _csv_file(contenido)}
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["rows"]


# ── perfil ───────────────────────────────────────────────────────────────────

def test_profile_se_crea_con_los_valores_por_defecto(client, db_session, auth_user):
    data = client.get("/me/profile").json()

    assert data["email"] == auth_user.email
    assert data["stock_minimo"] == 7
    assert data["dias_alerta_vencimiento"] == 7
    assert data["alert_check_hour"] == 5
    assert data["user_type"] == "individuo"
    assert db_session.query(UserProfile).filter(UserProfile.user_id == str(auth_user.id)).count() == 1


def test_profile_se_puede_actualizar_parcialmente(client, db_session, auth_user):
    client.get("/me/profile")  # crea el perfil

    resp = client.put("/me/profile", json={"first_name": "Gian", "stock_minimo": 3})

    assert resp.status_code == 200
    data = resp.json()
    assert data["first_name"] == "Gian"
    assert data["stock_minimo"] == 3
    assert data["dias_alerta_vencimiento"] == 7  # no tocado, sigue el default


def test_profile_username_y_apellido_tambien_se_actualizan(client, db_session, auth_user):
    resp = client.put(
        "/me/profile",
        json={"last_name": "Ubaldini", "username": "gianluuba", "dias_alerta_vencimiento": 15},
    )

    data = resp.json()
    assert data["last_name"] == "Ubaldini"
    assert data["username"] == "gianluuba"
    assert data["dias_alerta_vencimiento"] == 15


# ── stats ────────────────────────────────────────────────────────────────────

def test_my_stats_cuenta_bajo_stock_y_por_vencer(client, db_session, auth_user):
    bajo = create_product(db_session, auth_user, stock=1, categoria="Lácteos")
    create_stock_alert(db_session, auth_user, bajo, days_until_stockout=2)  # alcanza para menos que los 7 días del perfil
    create_product(db_session, auth_user, stock=50, vencimiento=ar_today(), categoria="Lácteos")
    create_product(db_session, auth_user, stock=50, vencimiento=date(2020, 1, 1), categoria="Vieja")
    create_product(db_session, auth_user, stock=50, imagen_url="http://x/1.jpg", categoria="Lácteos")

    data = client.get("/me/stats").json()

    assert data["total_products"] == 4
    assert data["total_units"] == 151
    assert data["products_with_image"] == 1
    assert data["low_stock_products"] == 1
    # vence hoy y ya venció: las dos tienen una alerta crítica de vencimiento, como en el inicio de la app
    assert data["expiring_soon_products"] == 2
    assert data["expired_products"] == 1
    assert data["critical_alerts"] == 3
    assert data["warning_alerts"] == 0
    labels = {c["label"] for c in data["products_by_category"]}
    assert "Lácteos" in labels


def test_my_stats_respeta_el_umbral_propio_del_producto(client, db_session, auth_user):
    lejos = create_product(db_session, auth_user, stock=5, stock_minimo=3)
    cerca = create_product(db_session, auth_user, stock=2, stock_minimo=3)
    create_stock_alert(db_session, auth_user, lejos, days_until_stockout=3)  # con su límite de 3 días no es crítica
    create_stock_alert(db_session, auth_user, cerca, days_until_stockout=2)  # alcanza para menos de 3 días

    data = client.get("/me/stats").json()

    assert data["low_stock_products"] == 1
    assert data["warning_alerts"] == 1


# ── listado ──────────────────────────────────────────────────────────────────

def test_my_products_filtra_por_busqueda_y_categoria(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Leche entera", marca="La Serenísima", categoria="Lácteos")
    create_product(db_session, auth_user, nombre="Detergente", marca="Ala", categoria="Limpieza")

    data = client.get("/me/products", params={"search": "serenísima"}).json()
    assert [p["nombre"] for p in data["items"]] == ["Leche entera"]

    data = client.get("/me/products", params={"categoria": "limpieza"}).json()
    assert [p["nombre"] for p in data["items"]] == ["Detergente"]


def test_my_products_filtra_por_vencido(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Vencido", vencimiento=date(2020, 1, 1))
    create_product(db_session, auth_user, nombre="Vigente", vencimiento=date(2030, 1, 1))
    create_product(db_session, auth_user, nombre="Sin fecha", vencimiento=None)

    vencidos = {p["nombre"] for p in client.get("/me/products", params={"vencido": True}).json()["items"]}
    assert vencidos == {"Vencido"}

    vigentes = {p["nombre"] for p in client.get("/me/products", params={"vencido": False}).json()["items"]}
    assert vigentes == {"Vigente", "Sin fecha"}


@pytest.mark.parametrize("campo", ["nombre", "stock", "vencimiento", "categoria"])
def test_my_products_se_puede_ordenar_por_cualquier_campo_sortable(client, db_session, auth_user, campo):
    create_product(db_session, auth_user, nombre="B", stock=1, categoria="B", vencimiento=date(2027, 1, 1))
    create_product(db_session, auth_user, nombre="A", stock=2, categoria="A", vencimiento=date(2026, 1, 1))

    data = client.get("/me/products", params={"sort_by": campo, "sort_dir": "desc"}).json()

    assert data["total"] == 2


def test_my_products_bajo_stock_y_por_vencer_se_resuelven_en_python(client, db_session, auth_user):
    bajo = create_product(db_session, auth_user, nombre="Bajo", stock=1)
    create_stock_alert(db_session, auth_user, bajo)
    create_product(db_session, auth_user, nombre="Normal", stock=99)
    por_vencer = create_product(db_session, auth_user, nombre="PorVencer", vencimiento=ar_today(), stock=99)
    create_product(db_session, auth_user, nombre="Vencido", vencimiento=date(2020, 1, 1), stock=99)  # ya venció: no es «por vencer»

    data = client.get("/me/products", params={"bajo_stock": True}).json()
    assert [p["id"] for p in data["items"]] == [bajo.id]

    data = client.get("/me/products", params={"por_vencer": True}).json()
    assert [p["id"] for p in data["items"]] == [por_vencer.id]


@pytest.mark.parametrize("campo", ["nombre", "stock", "vencimiento", "categoria", None])
def test_my_products_bajo_stock_se_puede_ordenar(client, db_session, auth_user, campo):
    for nombre, vencimiento in (("B", date(2027, 1, 1)), ("A", date(2026, 1, 1))):
        producto = create_product(db_session, auth_user, nombre=nombre, stock=1, categoria=nombre, vencimiento=vencimiento)
        create_stock_alert(db_session, auth_user, producto)

    params = {"bajo_stock": True, "sort_dir": "desc"}
    if campo:
        params["sort_by"] = campo
    data = client.get("/me/products", params=params).json()

    assert data["total"] == 2


# ── exportación ──────────────────────────────────────────────────────────────

def test_export_my_products_csv_con_filtros(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Leche", categoria="Lácteos")
    create_product(db_session, auth_user, nombre="Detergente", categoria="Limpieza")

    resp = client.get("/me/products/export", params={"format": "csv", "categoria": "Lácteos"})

    text = resp.content.decode("utf-8-sig")
    assert "Leche" in text
    assert "Detergente" not in text


def test_export_my_products_xlsx(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Leche")

    resp = client.get("/me/products/export", params={"format": "xlsx"})

    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")


def test_export_my_products_con_vencido_y_orden(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Vencido", vencimiento=date(2020, 1, 1))
    create_product(db_session, auth_user, nombre="Vigente", vencimiento=date(2030, 1, 1))

    resp = client.get(
        "/me/products/export",
        params={"format": "csv", "vencido": False, "sort_by": "nombre", "sort_dir": "desc"},
    )

    text = resp.content.decode("utf-8-sig")
    assert "Vigente" in text
    assert "Vencido" not in text


def test_export_my_products_con_bajo_stock(client, db_session, auth_user):
    create_stock_alert(db_session, auth_user, create_product(db_session, auth_user, nombre="Bajo", stock=1))
    create_product(db_session, auth_user, nombre="Alto", stock=99)

    resp = client.get("/me/products/export", params={"format": "csv", "bajo_stock": True})

    text = resp.content.decode("utf-8-sig")
    assert "Bajo" in text
    assert "Alto" not in text


def test_filter_options_lista_categorias_ubicaciones_y_marcas(client, db_session, auth_user):
    create_product(db_session, auth_user, categoria="Lácteos", ubicacion="Heladera", marca="La Serenísima")
    create_product(db_session, auth_user, categoria=None if False else "Otra", ubicacion=None, marca=None)

    data = client.get("/me/products/filters/options").json()

    assert "Lácteos" in data["categorias"]
    assert "Heladera" in data["ubicaciones"]
    assert "La Serenísima" in data["marcas"]


# ── defaults ─────────────────────────────────────────────────────────────────

def test_defaults_sin_perfil_devuelve_todo_vacio(client, auth_user):
    data = client.get("/me/products/defaults").json()

    assert data == {"last_categoria": None, "last_ubicacion": None}


def test_defaults_recuerda_la_ultima_categoria_y_ubicacion_usada(client, db_session, auth_user):
    client.post(
        "/me/products",
        json={"nombre": "Yerba", "categoria": "Almacén", "stock": 1, "ubicacion": "Alacena"},
    )

    data = client.get("/me/products/defaults").json()

    assert data == {"last_categoria": "Almacén", "last_ubicacion": "Alacena"}


# ── código de barras vacío ───────────────────────────────────────────────────

def test_by_barcode_con_separadores_solos_no_es_un_codigo(client, auth_user):
    resp = client.get("/me/products/by-barcode/--- ")

    assert resp.status_code == 200
    assert resp.json() is None


# ── edición en lote ──────────────────────────────────────────────────────────

def test_bulk_edit_sin_ids_es_400(client):
    resp = client.patch("/me/products/bulk-edit", json={"product_ids": []})

    assert resp.status_code == 400


def test_bulk_edit_sin_cambios_es_400(client, db_session, auth_user):
    p = create_product(db_session, auth_user)

    resp = client.patch("/me/products/bulk-edit", json={"product_ids": [p.id]})

    assert resp.status_code == 400


def test_bulk_edit_de_productos_inexistentes_es_404(client):
    resp = client.patch("/me/products/bulk-edit", json={"product_ids": [999999], "categoria": "X"})

    assert resp.status_code == 404


def test_bulk_edit_actualiza_y_audita(client, db_session, auth_user):
    p1 = create_product(db_session, auth_user, categoria="Vieja")
    p2 = create_product(db_session, auth_user, categoria="Vieja")

    resp = client.patch(
        "/me/products/bulk-edit",
        json={"product_ids": [p1.id, p2.id], "categoria": "Nueva", "ubicacion": "Alacena", "marca": "Genérica"},
    )

    assert resp.status_code == 200
    assert resp.json()["updated"] == 2
    db_session.refresh(p1)
    db_session.refresh(p2)
    assert (p1.categoria, p1.ubicacion, p1.marca) == ("Nueva", "Alacena", "Genérica")
    log = db_session.query(AdminActivityLog).filter(AdminActivityLog.event_type == "product_updated").one()
    assert "2 productos" in log.description


# ── eventos de consumo: casos borde ───────────────────────────────────────────

def test_corregir_un_evento_inexistente_da_404(client, db_session, auth_user):
    p = create_product(db_session, auth_user, stock=10)

    resp = client.put(f"/me/products/{p.id}/events/999999", json={"quantity": 1, "event_type": "consumed"})

    assert resp.status_code == 404


def test_borrar_un_evento_que_dejaria_stock_negativo_es_422(client, db_session, auth_user):
    p = create_product(db_session, auth_user, stock=5)
    # Entra stock de más (restocked) y después se consume todo, incluido lo recién entrado.
    client.patch(f"/me/products/{p.id}", json={"stock": 15})  # restocked +10
    client.post("/me/products/salida", json={"product_id": p.id, "quantity": 12, "event_type": "consumed"})
    restock_event = db_session.query(ConsumptionEvent).filter(
        ConsumptionEvent.product_id == p.id, ConsumptionEvent.event_type == "restocked"
    ).one()

    resp = client.delete(f"/me/products/{p.id}/events/{restock_event.id}")

    assert resp.status_code == 422
    db_session.refresh(p)
    assert p.stock == 3


# ── bulk parse: validaciones de fila ──────────────────────────────────────────

def test_bulk_parse_nombre_obligatorio(client):
    filas = _parse(client, "nombre,categoria,stock\n,Farmacia,3\n")

    assert filas[0]["valid"] is False
    assert "nombre es obligatorio" in filas[0]["errors"]


def test_bulk_parse_categoria_obligatoria_solo_en_modo_entrada(client):
    filas = _parse(client, "nombre,categoria,stock\nIbuprofeno,,3\n", mode="entrada")
    assert "categoria es obligatoria" in filas[0]["errors"]

    filas = _parse(client, "nombre,categoria,stock\nIbuprofeno,,3\n", mode="salida")
    assert "categoria es obligatoria" not in filas[0]["errors"]


@pytest.mark.parametrize(
    "stock, mensaje",
    [
        ("", "stock es obligatorio"),
        ("no-es-numero", "stock debe ser un número entero"),
        ("0", "stock debe ser mayor a 0"),
        ("-5", "stock debe ser mayor a 0"),
    ],
)
def test_bulk_parse_valida_el_stock(client, stock, mensaje):
    filas = _parse(client, f"nombre,categoria,stock\nIbuprofeno,Farmacia,{stock}\n")

    assert mensaje in filas[0]["errors"]


def test_bulk_parse_vencimiento_invalido(client):
    filas = _parse(client, "nombre,categoria,stock,vencimiento\nIbuprofeno,Farmacia,1,no-es-fecha\n")

    assert any("vencimiento" in e for e in filas[0]["errors"])


def test_bulk_parse_vencimiento_formato_dd_mm_aaaa(client):
    filas = _parse(client, "nombre,categoria,stock,vencimiento\nIbuprofeno,Farmacia,1,31/12/2026\n")

    assert filas[0]["valid"] is True
    assert filas[0]["data"]["vencimiento"] == "2026-12-31"


@pytest.mark.parametrize(
    "columna",
    ["stock_minimo", "dias_alerta_vencimiento"],
)
@pytest.mark.parametrize("valor, mensaje_parte", [("abc", "debe ser un número entero"), ("0", "debe ser mayor a 0")])
def test_bulk_parse_valida_los_enteros_opcionales(client, columna, valor, mensaje_parte):
    filas = _parse(client, f"nombre,categoria,stock,{columna}\nIbuprofeno,Farmacia,1,{valor}\n")

    assert any(mensaje_parte in e for e in filas[0]["errors"])


def test_bulk_parse_formato_no_soportado(client):
    resp = client.post(
        "/me/products/bulk/parse",
        files={"file": ("productos.txt", b"nombre,categoria,stock\n", "text/plain")},
    )

    assert resp.status_code == 400


# ── bulk parse: entrada (merge dentro del mismo archivo) ──────────────────────

def test_bulk_parse_entrada_junta_dos_filas_del_mismo_producto_en_el_archivo(client):
    filas = _parse(
        client,
        "nombre,categoria,stock\nLeche,Lácteos,2\nLeche,Lácteos,3\n",
    )

    assert filas[0]["action"] == "create"
    assert filas[1]["action"] == "merge"
    assert filas[1]["existing_stock"] == 2


def test_bulk_parse_entrada_detecta_un_producto_que_ya_existe(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Leche", marca=None, categoria="Lácteos", stock=7, vencimiento=None)

    filas = _parse(client, "nombre,categoria,stock\nLeche,Lácteos,2\n")

    assert filas[0]["action"] == "merge"
    assert filas[0]["existing_stock"] == 7


# ── bulk parse: salida ─────────────────────────────────────────────────────────

def test_bulk_parse_salida_sin_producto_existente_es_invalida(client):
    filas = _parse(client, "nombre,categoria,stock\nFantasma,Farmacia,1\n", mode="salida")

    assert filas[0]["valid"] is False
    assert "no existe un producto tuyo" in filas[0]["errors"][0]


def test_bulk_parse_salida_sin_stock_suficiente_es_invalida(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Ibuprofeno", marca=None, categoria="Farmacia", stock=2, vencimiento=None)

    filas = _parse(client, "nombre,categoria,stock\nIbuprofeno,Farmacia,5\n", mode="salida")

    assert filas[0]["valid"] is False
    assert "no hay suficiente stock disponible" in filas[0]["errors"][0]


def test_bulk_parse_salida_descuenta_localmente_entre_filas(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Ibuprofeno", marca=None, categoria="Farmacia", stock=5, vencimiento=None)

    filas = _parse(
        client,
        "nombre,categoria,stock\nIbuprofeno,Farmacia,3\nIbuprofeno,Farmacia,3\n",
        mode="salida",
    )

    assert filas[0]["valid"] is True
    assert filas[1]["valid"] is False  # ya no queda stock suficiente (5 - 3 = 2 < 3)


# ── bulk parse: xlsx ─────────────────────────────────────────────────────────

def _xlsx_bytes(filas: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["nombre", "categoria", "stock", "vencimiento"])
    for fila in filas:
        ws.append(fila)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_bulk_parse_lee_xlsx_con_fecha_como_objeto_date(client):
    contenido = _xlsx_bytes([["Ibuprofeno", "Farmacia", 3, date(2026, 12, 31)]])

    resp = client.post(
        "/me/products/bulk/parse",
        params={"mode": "entrada"},
        files={"file": ("productos.xlsx", contenido,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    filas = resp.json()["rows"]
    assert filas[0]["valid"] is True
    assert filas[0]["data"]["vencimiento"] == "2026-12-31"


def test_bulk_parse_lee_xlsx_con_fecha_como_datetime(client):
    contenido = _xlsx_bytes([["Ibuprofeno", "Farmacia", 3, datetime(2026, 12, 31, 10, 30)]])

    resp = client.post(
        "/me/products/bulk/parse",
        params={"mode": "entrada"},
        files={"file": ("productos.xlsx", contenido,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    filas = resp.json()["rows"]
    assert filas[0]["data"]["vencimiento"] == "2026-12-31"


def test_bulk_parse_xlsx_vacio_no_tiene_filas(client):
    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)

    resp = client.post(
        "/me/products/bulk/parse",
        files={"file": ("vacio.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert resp.json()["rows"] == []


def test_bulk_parse_xlsx_solo_encabezado_no_tiene_filas(client):
    wb = Workbook()
    ws = wb.active
    ws.append(["nombre", "categoria", "stock"])
    buf = io.BytesIO()
    wb.save(buf)

    resp = client.post(
        "/me/products/bulk/parse",
        files={"file": ("solo_encabezado.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert resp.json()["rows"] == []


# ── bulk commit ────────────────────────────────────────────────────────────────

def test_bulk_commit_entrada_crea_un_producto_nuevo(client, db_session, auth_user):
    resp = client.post("/me/products/bulk/commit", json={
        "mode": "entrada", "origen": "Prueba",
        "rows": [{"row_number": 1, "nombre": "Yerba", "categoria": "Almacén", "stock": 5}],
    })

    assert resp.status_code == 200
    body = resp.json()
    assert (body["created"], body["merged"], body["failed"]) == (1, 0, 0)
    producto = db_session.query(Product).filter(Product.nombre == "Yerba").one()
    assert producto.stock == 5
    log = db_session.query(AdminActivityLog).filter(AdminActivityLog.event_type == "product_created").one()
    assert "cargó" in log.description


def test_bulk_commit_entrada_suma_a_un_producto_existente_y_completa_el_codigo(client, db_session, auth_user):
    producto = create_product(db_session, auth_user, nombre="Yerba", marca=None, categoria="Almacén",
                              stock=5, vencimiento=None, codigo_barras=None)

    resp = client.post("/me/products/bulk/commit", json={
        "mode": "entrada", "origen": "Prueba",
        "rows": [{"row_number": 1, "nombre": "Yerba", "categoria": "Almacén", "stock": 3,
                  "codigo_barras": "7794640172021"}],
    })

    body = resp.json()
    assert (body["created"], body["merged"]) == (0, 1)
    db_session.refresh(producto)
    assert producto.stock == 8
    assert producto.codigo_barras == "7794640172021"


def test_bulk_commit_salida_descuenta_stock(client, db_session, auth_user):
    producto = create_product(db_session, auth_user, nombre="Yerba", marca=None, categoria="Almacén",
                              stock=10, vencimiento=None)

    resp = client.post("/me/products/bulk/commit", json={
        "mode": "salida", "origen": "Documento Comercial",
        "rows": [{"row_number": 1, "nombre": "Yerba", "categoria": "Almacén", "stock": 4}],
    })

    body = resp.json()
    assert (body["merged"], body["failed"]) == (1, 0)
    db_session.refresh(producto)
    assert producto.stock == 6
    evento = db_session.query(ConsumptionEvent).filter(ConsumptionEvent.product_id == producto.id).one()
    assert evento.event_type == "consumed"


def test_bulk_commit_salida_de_un_producto_inexistente_falla(client):
    resp = client.post("/me/products/bulk/commit", json={
        "mode": "salida", "origen": "X",
        "rows": [{"row_number": 1, "nombre": "Fantasma", "categoria": "X", "stock": 1}],
    })

    body = resp.json()
    assert body["failed"] == 1
    assert "no existe un producto tuyo" in body["errors"][0]


def test_bulk_commit_salida_sin_stock_suficiente_falla(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Yerba", marca=None, categoria="Almacén", stock=2, vencimiento=None)

    resp = client.post("/me/products/bulk/commit", json={
        "mode": "salida", "origen": "X",
        "rows": [{"row_number": 1, "nombre": "Yerba", "categoria": "Almacén", "stock": 5}],
    })

    body = resp.json()
    assert body["failed"] == 1
    assert "no hay suficiente stock disponible" in body["errors"][0]


def test_bulk_commit_un_codigo_invalido_no_frena_al_resto(client, db_session, auth_user):
    # La fila que falla va última: el `db.rollback()` que hace el endpoint ante un
    # código inválido expira toda la sesión (incluida `current`, ya cargada desde
    # antes de la petición) y una fila exitosa *después* de esa reversión
    # necesitaría releerla — un artefacto del savepoint de test, no del código bajo
    # prueba. Alcanza con probar que una fila previa exitosa no se pierde.
    resp = client.post("/me/products/bulk/commit", json={
        "mode": "entrada", "origen": "X",
        "rows": [
            {"row_number": 1, "nombre": "Bueno", "categoria": "X", "stock": 1},
            {"row_number": 2, "nombre": "Malo", "categoria": "X", "stock": 1, "codigo_barras": "12345"},
        ],
    })

    body = resp.json()
    # La fila 2 no le pisa el resultado a la fila 1: cada una se cuenta por su cuenta.
    assert body["failed"] == 1
    assert body["created"] == 1
    assert body["errors"][0].startswith("Fila 2:")


def test_bulk_commit_un_error_inesperado_se_informa_sin_frenar_el_resto(client, db_session, auth_user, monkeypatch):
    # Cubre la rama de `except Exception` (no `HTTPException`) del commit: se
    # fuerza un error genérico al normalizar el código de la fila que tiene que
    # fallar, sin tocar `Product` (reemplazar la clase completa rompería también
    # las consultas de `_find_matching_product`, que la usan para armar el WHERE).
    original_normalizar = me_api._normalizar_codigo

    def normalizar_que_a_veces_revienta(codigo):
        if codigo == "BOOM":
            raise SQLAlchemyError("boom")
        return original_normalizar(codigo)

    monkeypatch.setattr(me_api, "_normalizar_codigo", normalizar_que_a_veces_revienta)

    # La fila que falla va última por la misma razón que en el test anterior: el
    # `db.rollback()` del `except Exception` expira la sesión completa.
    resp = client.post("/me/products/bulk/commit", json={
        "mode": "entrada", "origen": "X",
        "rows": [
            {"row_number": 1, "nombre": "Bien", "categoria": "X", "stock": 1},
            {"row_number": 2, "nombre": "Falla", "categoria": "X", "stock": 1, "codigo_barras": "BOOM"},
        ],
    })

    body = resp.json()
    assert body["failed"] == 1
    assert body["created"] == 1
    assert "Fila 2: boom" in body["errors"][0]


# ── orden alfabético y sin fecha de vencimiento ──────────────────────────────

def mis_nombres(client, **params):
    return [p["nombre"] for p in client.get("/me/products", params=params).json()["items"]]


def test_my_products_ordena_por_nombre_sin_distinguir_mayusculas_ni_tildes(client, db_session, auth_user):
    for nombre in ["zanahoria", "Ñoquis", "Álvaro", "banana", "Arroz"]:
        create_product(db_session, auth_user, nombre=nombre)

    assert mis_nombres(client, sort_by="nombre") == ["Álvaro", "Arroz", "banana", "Ñoquis", "zanahoria"]
    assert mis_nombres(client, sort_by="nombre", sort_dir="desc") == ["zanahoria", "Ñoquis", "banana", "Arroz", "Álvaro"]


def test_my_products_por_vencimiento_deja_afuera_los_que_no_vencen(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Sin fecha")
    create_product(db_session, auth_user, nombre="Tarde", vencimiento=date(2027, 5, 1))
    create_product(db_session, auth_user, nombre="Pronto", vencimiento=date(2027, 1, 1))

    data = client.get("/me/products", params={"sort_by": "vencimiento"}).json()

    assert [p["nombre"] for p in data["items"]] == ["Pronto", "Tarde"]
    assert data["total"] == 2


def test_my_products_las_paginas_no_repiten_productos(client, db_session, auth_user):
    for i in range(7):
        create_product(db_session, auth_user, nombre=f"P{i}", categoria="Igual", stock=1)

    primera = client.get("/me/products", params={"sort_by": "stock", "limit": 4}).json()["items"]
    segunda = client.get("/me/products", params={"sort_by": "stock", "limit": 4, "offset": 4}).json()["items"]

    assert len({p["id"] for p in primera + segunda}) == 7


def test_my_products_con_bajo_stock_tambien_ordena_alfabeticamente_y_sin_los_que_no_vencen(client, db_session, auth_user):
    for nombre, vencimiento in (("banana", date(2027, 2, 1)), ("Álvaro", date(2027, 1, 1)), ("Sin fecha", None)):
        create_stock_alert(db_session, auth_user, create_product(db_session, auth_user, nombre=nombre, stock=1, vencimiento=vencimiento))

    assert mis_nombres(client, bajo_stock=True, sort_by="nombre") == ["Álvaro", "banana", "Sin fecha"]
    data = client.get("/me/products", params={"bajo_stock": True, "sort_by": "vencimiento"}).json()
    assert [p["nombre"] for p in data["items"]] == ["Álvaro", "banana"]
    assert data["total"] == 2


# ── ubicación y stock total del inventario ───────────────────────────────────

def test_my_products_filtra_por_ubicacion_sin_distinguir_mayusculas(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Leche", ubicacion="Heladera")
    create_product(db_session, auth_user, nombre="Arroz", ubicacion="Alacena")

    data = client.get("/me/products", params={"ubicacion": "heladera"}).json()

    assert [p["nombre"] for p in data["items"]] == ["Leche"]


def test_el_export_tambien_filtra_por_ubicacion(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Leche", ubicacion="Heladera")
    create_product(db_session, auth_user, nombre="Arroz", ubicacion="Alacena")

    body = client.get("/me/products/export", params={"ubicacion": "Alacena"}).text

    assert "Arroz" in body and "Leche" not in body


def test_el_stock_total_suma_los_productos_de_la_lista_con_sus_filtros(client, db_session, auth_user):
    create_product(db_session, auth_user, nombre="Leche", categoria="Lácteos", ubicacion="Heladera", stock=4)
    create_product(db_session, auth_user, nombre="Yogur", categoria="Lácteos", ubicacion="Alacena", stock=6)
    create_product(db_session, auth_user, nombre="Arroz", categoria="Almacén", ubicacion="Alacena", stock=10, vencimiento=date(2020, 1, 1))
    create_product(db_session, auth_user, nombre="Ajeno", stock=99, user_id="otro-usuario")

    assert client.get("/me/products/stock-total").json() == {"stock_total": 20, "products": 3}
    assert client.get("/me/products/stock-total", params={"categoria": "lácteos"}).json() == {"stock_total": 10, "products": 2}
    assert client.get("/me/products/stock-total", params={"ubicacion": "alacena", "search": "yog"}).json() == {"stock_total": 6, "products": 1}
    assert client.get("/me/products/stock-total", params={"vencido": True}).json() == {"stock_total": 10, "products": 1}


def test_el_stock_total_respeta_el_filtro_de_alertas(client, db_session, auth_user, monkeypatch):
    keep = create_product(db_session, auth_user, nombre="Se acaba", stock=2)
    create_product(db_session, auth_user, nombre="Sobra", stock=50)
    monkeypatch.setattr(me_api, "_alert_filter_ids", lambda db, uid, bajo_stock, por_vencer: {keep.id})

    assert client.get("/me/products/stock-total", params={"bajo_stock": True}).json() == {"stock_total": 2, "products": 1}
